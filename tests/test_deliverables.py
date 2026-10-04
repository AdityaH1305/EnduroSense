"""The report and slide builders: they run, and what they say matches the result files."""
import importlib.util
import json
import re
import sys

import pandas as pd
import pytest

from endurosense.config import ROOT, data_path

sys.path.insert(0, str(ROOT / "app"))
import logic as L  # noqa: E402  (the facts the deliverables are built from)

FINAL = data_path("results") / "final"
pytestmark = [pytest.mark.data, pytest.mark.skipif(not (FINAL / "success_criteria.csv").exists(), reason="run scripts/run_all.py --final first")]


def load(name):
    spec = importlib.util.spec_from_file_location(name.replace(".py", ""), ROOT / "docs" / "build" / name)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_report_builds_and_quotes_the_final_results(tmp_path):
    from docx import Document
    report = load("build_report.py")
    out = tmp_path / "report.docx"
    report.build(out)
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    a = pd.read_csv(FINAL / "model_a_test_metrics.csv").set_index("model")
    pts = pd.read_csv(FINAL / "operating_points.csv")
    pts = pts[(pts["pairs"] == "all pairs") & (pts["margins"] == "tuned on development data")].set_index("policy")["unsafe_approval_rate"]
    b = pd.read_csv(FINAL / "model_b_test_metrics.csv")
    b = b[b["model"] == "Physics-first"].set_index("flights_group")["mape_pct"]
    for expected in (f"{a.loc['GRU ensemble, calibrated (main)', 'test_mae_supported']:.2f} Wh",
                     f"{a.loc['Voltage lookup', 'test_mae_supported']:.2f} Wh",
                     f"{100 * pts['P1 minutes left, as in the brief (ratio >= 1)']:.1f}%", f"{100 * pts['P3 EnduroSense, tau = 0.95']:.2f}%",
                     f"{b['unseen routes']:.1f}%", "Not met", "Not met, within sampling noise"):
        assert expected in text, expected
    crit = pd.read_csv(FINAL / "success_criteria.csv")
    met_in_report = [row.cells[1].text for row in doc.tables[0].rows[1:]]
    assert [m == "Met" for m in met_in_report] == crit["met"].tolist()          # the verdicts are the evaluation's, not the author's
    assert len(doc.tables) == 5 and len(doc.inline_shapes) == 4                    # five tables, four figures
    fleet = pd.read_csv(FINAL / "fleet_simulation.csv").set_index("policy")["unsafe_per_100_missions"]
    cells = [row.cells[1].text for row in doc.tables[4].rows[1:]]                  # unsafe missions per 100: one decimal above 1, two below
    assert cells[0] == f"{fleet['P1 minutes left (brief)']:.1f}" and cells[3] == f"{fleet['P3 EnduroSense (tau = 0.95)']:.2f}"
    facts = L.headline_facts()
    assert [row.cells[1].text for row in doc.tables[0].rows[1:]] == facts["verdicts"]
    assert f"{facts['number_word'](facts['verification_passes']).capitalize()} verification passes" in text
    assert f"{facts['a_gru_cv']:.2f} Wh" in text and "judged once" not in text
    assert "**" not in text and "`" not in text                                   # no stray markup


def test_slides_specification_is_complete_and_within_the_slide(tmp_path, monkeypatch):
    slides = load("build_slides.py")
    monkeypatch.setattr(slides, "FIG", tmp_path)                                  # charts go to a temporary folder
    deck = slides.deck()
    assert len(deck) == 13 and deck[0]["title"] == "" and all(s["title"] for s in deck[1:]) and all(s["notes"] for s in deck)
    for s in deck:
        for e in s["elements"]:
            assert e["x"] >= 0 and e["y"] >= 0 and e["x"] + e["w"] <= slides.W_IN + 0.01, (s["title"], e.get("text", e.get("path")))
            assert e["y"] + e["h"] <= slides.H_IN + 0.01, (s["title"], e.get("text", e.get("path")))
            if e["type"] == "text":
                assert "**" not in e["text"]
                for start, length in e["bold_ranges"]:                             # bold spans lie inside the text
                    assert 1 <= start and start + length - 1 <= len(e["text"])
            if e["type"] == "image":
                assert (tmp_path / e["path"].replace("\\", "/").split("/")[-1]).exists()
    everything = json.dumps(deck)
    pts = pd.read_csv(FINAL / "operating_points.csv")
    p3 = pts[(pts["pairs"] == "all pairs") & (pts["margins"] == "tuned on development data") & (pts["policy"] == "P3 EnduroSense, tau = 0.95")]
    assert f"{100 * p3['unsafe_approval_rate'].iloc[0]:.2f}%" in everything
    boxes = [e["text"] for e in deck[7]["elements"] if e["type"] == "box"]
    assert boxes == ["GO: EnduroSense", "NO-GO: EnduroSense", "GO: EnduroSense"]   # the middle mission is the one refused
    lead = next(e for s in deck for e in s["elements"] if e["type"] == "text" and e["text"].startswith("The brief:"))
    bold = [lead["text"][a - 1:a - 1 + n] for a, n in lead["bold_ranges"]]          # positions are 1-based, as PowerPoint counts
    assert bold == ["The brief:", "The trouble with minutes:", "A single number hides how sure it is."]
    assert re.search(r"1 in \d+ missions", everything) and "verification passes" in everything
